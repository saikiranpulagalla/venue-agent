from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import MSO_AUTO_SIZE

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "backend" / "tests" / "fixtures"
OUT.mkdir(parents=True, exist_ok=True)


def add_text(slide, text, left, top, width, height, size):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.clear()
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.name = "DejaVu Sans"
    return box


def fixture(name, body_size=26, footer_size=None, image_only=False):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    if not image_only:
        add_text(slide, "Venue Agent Demo", 0.8, 0.6, 11.5, 0.8, 34)
        add_text(slide, "Room-aware structural analysis sample text.", 0.9, 2.0, 11.0, 1.2, body_size)
        if footer_size:
            add_text(slide, "Project-owned fixture — WCC 2026", 0.9, 6.8, 5.0, 0.3, footer_size)
    else:
        # no structural text by design
        shape = slide.shapes.add_shape(1, Inches(1), Inches(1), Inches(5), Inches(3))
        shape.fill.solid()
    prs.save(OUT / name)


fixture("safe_large.pptx", body_size=32)
fixture("small_text.pptx", body_size=12)
fixture("mixed_sizes.pptx", body_size=24, footer_size=8)
fixture("image_only.pptx", image_only=True)
