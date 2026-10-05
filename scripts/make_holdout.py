from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import MSO_AUTO_SIZE
import json
ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'validation'/'holdout'; OUT.mkdir(parents=True,exist_ok=True)

def deck(name,size):
    prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
    sl=prs.slides.add_slide(prs.slide_layouts[6]); box=sl.shapes.add_textbox(Inches(1),Inches(2),Inches(10),Inches(1.5))
    r=box.text_frame.paragraphs[0].add_run(); r.text=f'Holdout {name} structural text'; r.font.size=Pt(size); r.font.name='DejaVu Sans'; prs.save(OUT/name)

deck('obvious_large.pptx',44)
deck('obvious_tiny.pptx',7)
(OUT/'expected.json').write_text(json.dumps({
  'venue': {'active_image_height_m':1.8,'farthest_viewer_distance_m':10.8,'measurement_basis':'MEASURED'},
  'obvious_large.pptx':'MEETS_TARGET',
  'obvious_tiny.pptx':'BELOW_TARGET'
},indent=2))
