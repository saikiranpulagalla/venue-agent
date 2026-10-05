from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import MSO_AUTO_SIZE

from app.domain.models import AnalysisCapability, MappingConfidence, MutationCapability
from app.pptx.parser import map_elements
from app.pptx.renderer import render_pptx_to_pdf

pytestmark = pytest.mark.resource


def test_same_text_in_two_shapes_keeps_distinct_rendered_identity(tmp_path: Path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    left = slide.shapes.add_textbox(Inches(0.8), Inches(1.2), Inches(4.2), Inches(1.0))
    left.text_frame.auto_size = MSO_AUTO_SIZE.NONE
    left.text_frame.paragraphs[0].text = "Repeated label"
    left.text_frame.paragraphs[0].runs[0].font.name = "DejaVu Sans"
    left.text_frame.paragraphs[0].runs[0].font.size = Pt(12)

    right = slide.shapes.add_textbox(Inches(5.5), Inches(4.5), Inches(3.8), Inches(1.0))
    right.text_frame.auto_size = MSO_AUTO_SIZE.NONE
    right.text_frame.paragraphs[0].text = "Repeated label"
    right.text_frame.paragraphs[0].runs[0].font.name = "DejaVu Sans"
    right.text_frame.paragraphs[0].runs[0].font.size = Pt(30)

    source = tmp_path / "duplicate-text.pptx"
    prs.save(source)
    pdf = render_pptx_to_pdf(source, tmp_path / "render")
    elements = map_elements(source, pdf)

    assert len(elements) == 2
    by_shape = {e.shape_id: e for e in elements}
    small = by_shape[left.shape_id]
    large = by_shape[right.shape_id]

    assert small.mapping_confidence in {MappingConfidence.EXACT, MappingConfidence.STRONG}
    assert large.mapping_confidence in {MappingConfidence.EXACT, MappingConfidence.STRONG}
    assert small.analysis_capability == AnalysisCapability.ANALYZABLE
    assert large.analysis_capability == AnalysisCapability.ANALYZABLE
    assert small.mutation_capability == MutationCapability.SAFE_MUTATION
    assert large.mutation_capability == MutationCapability.SAFE_MUTATION
    assert small.rendered_height_percent is not None
    assert large.rendered_height_percent is not None
    assert large.rendered_height_percent > small.rendered_height_percent * 1.8
