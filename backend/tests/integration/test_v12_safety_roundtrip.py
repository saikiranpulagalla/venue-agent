from pathlib import Path
import tempfile

import pytest
from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.util import Inches, Pt

from app.domain.models import MutationCapability, VenueProfile, WorkflowState
from app.pptx.mutations import apply_scale_text
from app.repair.verifier import verify_output
from app.venue.analyzer import VenueAnalyzer
from app.venue.reference import PublicBDMReference

pytestmark = pytest.mark.resource


def _analyzer(project_root):
    return VenueAnalyzer(PublicBDMReference(project_root / 'reference' / 'PUBLIC_BDM_REFERENCE_V1.json'))


def _add_box(slide, text, left, top, width, height, size, font='DejaVu Sans', wrap=True):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    box.text_frame.auto_size = MSO_AUTO_SIZE.NONE
    box.text_frame.word_wrap = wrap
    box.text_frame.clear()
    r = box.text_frame.paragraphs[0].add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.name = font
    return box


def _add_opaque_rectangle(slide, left, top, width, height):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(0, 0, 0)
    shape.line.fill.background()
    return shape


def test_visual_overlap_or_wrap_cannot_receive_verified(project_root):
    with tempfile.TemporaryDirectory() as td:
        td = Path(td)
        src = td / 'overlap-source.pptx'
        out = td / 'overlap-output.pptx'
        prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        target_shape = _add_box(
            slide,
            'Critical venue instruction with several words',
            1, 1.0, 5.0, 0.8, 12,
        )
        _add_box(
            slide,
            'SECOND LINE SHOULD NOT BE OVERLAPPED',
            1, 1.55, 7.0, 0.7, 16,
        )
        prs.save(src)

        venue = VenueProfile(active_image_height_m=1.8, farthest_viewer_distance_m=10.8, measurement_basis='MEASURED')
        analyzer = _analyzer(project_root)
        baseline = analyzer.analyze(src, venue, td / 'baseline')
        target = next(r for r in baseline.results if r.element.shape_id == target_shape.shape_id)
        apply_scale_text(src, out, target.element.slide_index, target.element.shape_id, 1.60)
        report = verify_output(src, out, baseline, target.element.element_id, analyzer, td / 'verify')

        assert report.final_state == WorkflowState.REVIEW_REQUIRED
        assert report.layout_safe is False
        assert any(
            reason.startswith('target_line_count_changed:') or reason.startswith('new_or_increased_text_overlap:')
            for reason in report.reasons
        ), report.model_dump()


def test_missing_requested_font_is_never_auto_mutable(project_root):
    with tempfile.TemporaryDirectory() as td:
        td = Path(td)
        src = td / 'missing-font.pptx'
        prs = Presentation(); slide = prs.slides.add_slide(prs.slide_layouts[6])
        _add_box(slide, 'Missing font geometry', 1, 1, 7, 1, 18, font='DefinitelyMissingVenueFont')
        prs.save(src)

        result = _analyzer(project_root).analyze(
            src,
            VenueProfile(active_image_height_m=1.8, farthest_viewer_distance_m=10.8, measurement_basis='MEASURED'),
            td / 'analysis',
        )
        assert result.results
        e = result.results[0].element
        assert e.mutation_capability == MutationCapability.REVIEW_ONLY
        assert 'font_substitution_or_unverified' in e.notes
        assert e.requested_font_families == ['DefinitelyMissingVenueFont']
        assert e.rendered_font_families


def test_table_text_remains_explicitly_not_analyzed(project_root):
    with tempfile.TemporaryDirectory() as td:
        td = Path(td)
        src = td / 'table.pptx'
        prs = Presentation(); slide = prs.slides.add_slide(prs.slide_layouts[6])
        table = slide.shapes.add_table(1, 1, Inches(1), Inches(1), Inches(5), Inches(1)).table
        table.cell(0, 0).text = 'Visible table text that must not disappear from coverage'
        prs.save(src)

        result = _analyzer(project_root).analyze(
            src,
            VenueProfile(active_image_height_m=1.8, farthest_viewer_distance_m=10.8, measurement_basis='MEASURED'),
            td / 'analysis',
        )
        assert result.summary.total_detected_text_elements == 1
        assert result.summary.not_analyzed_elements == 1
        assert result.summary.unsupported_visible_content_elements == 1
        assert result.summary.meets_target_elements == 0


def test_candidate_simulation_blocks_target_seeking_scale_that_wraps_or_overlaps(project_root):
    from app.repair.candidates import generate_candidates
    from app.repair.simulation import simulate_candidate

    with tempfile.TemporaryDirectory() as td:
        td = Path(td)
        src = td / 'candidate-overlap.pptx'
        prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        target_shape = _add_box(
            slide,
            'Critical venue instruction with several words',
            1, 1.0, 5.0, 0.8, 12,
        )
        _add_box(slide, 'SECOND LINE SHOULD NOT BE OVERLAPPED', 1, 1.55, 7.0, 0.7, 16)
        prs.save(src)

        venue = VenueProfile(active_image_height_m=1.8, farthest_viewer_distance_m=10.8, measurement_basis='MEASURED')
        analyzer = _analyzer(project_root)
        baseline = analyzer.analyze(src, venue, td / 'baseline')
        target_id = next(r.element.element_id for r in baseline.results if r.element.shape_id == target_shape.shape_id)
        raw = [c for c in generate_candidates(baseline) if c.issue_element_id == target_id]
        simulated = [simulate_candidate(src, baseline, c, analyzer) for c in raw]
        target_seeking = [c for c in simulated if c.predicted_state and c.predicted_state.value == 'MEETS_TARGET']

        assert target_seeking, [c.model_dump() for c in simulated]
        assert all(c.safe is False for c in target_seeking)
        assert any(
            any(r.startswith('target_line_count_changed:') or r.startswith('new_or_increased_text_overlap:') for r in c.policy_reasons)
            for c in target_seeking
        )


def test_unknown_content_remaining_after_successful_target_repair_cannot_be_verified(project_root):
    with tempfile.TemporaryDirectory() as td:
        td = Path(td)
        src = td / 'known-plus-unknown.pptx'
        out = td / 'known-plus-unknown-output.pptx'
        prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        target_shape = _add_box(slide, 'Small supported text', 1, 1, 8, 1, 12)
        table = slide.shapes.add_table(1, 1, Inches(1), Inches(4), Inches(5), Inches(1)).table
        table.cell(0, 0).text = 'Unsupported visible table text'
        prs.save(src)

        venue = VenueProfile(active_image_height_m=1.8, farthest_viewer_distance_m=10.8, measurement_basis='MEASURED')
        analyzer = _analyzer(project_root)
        baseline = analyzer.analyze(src, venue, td / 'baseline')
        target = next(r for r in baseline.results if r.element.shape_id == target_shape.shape_id)
        assert baseline.summary.unsupported_visible_content_elements == 1

        apply_scale_text(src, out, target.element.slide_index, target.element.shape_id, 1.60)
        report = verify_output(src, out, baseline, target.element.element_id, analyzer, td / 'verify')

        assert report.target_improved is True
        assert report.analysis_universe_preserved is True
        assert report.layout_safe is True
        assert report.no_unknown_or_uncovered is False
        assert report.final_state == WorkflowState.REVIEW_REQUIRED
        assert 'unknown_or_uncovered_content_remains' in report.reasons


@pytest.mark.parametrize("visual_kind", ["chart", "picture"])
def test_unsupported_visible_content_can_never_be_silently_verified(project_root, visual_kind):
    """A repairable text box does not erase coverage gaps from other visuals."""
    with tempfile.TemporaryDirectory() as td:
        td = Path(td)
        src = td / f"known-plus-{visual_kind}.pptx"
        out = td / f"known-plus-{visual_kind}-output.pptx"
        prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        target_shape = _add_box(slide, "Small supported text", 1, 1, 8, 1, 12)
        if visual_kind == "chart":
            data = ChartData(); data.categories = ["A", "B"]; data.add_series("Critical labels", (1, 2))
            slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(4), Inches(5), Inches(2), data)
        else:
            # A project-owned one-pixel PNG. Its contents are irrelevant: the
            # policy is that an overlapping arbitrary picture is outside
            # structural coverage.
            image = td / "decorative.png"
            image.write_bytes(bytes.fromhex("89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c4890000000d49444154789c6360f8cfc0000003010100c9fe92ef0000000049454e44ae426082"))
            slide.shapes.add_picture(str(image), Inches(1), Inches(1), Inches(5), Inches(1))
        prs.save(src)

        analyzer = _analyzer(project_root)
        venue = VenueProfile(active_image_height_m=1.8, farthest_viewer_distance_m=10.8, measurement_basis="MEASURED")
        baseline = analyzer.analyze(src, venue, td / "baseline")
        target = next(r for r in baseline.results if r.element.shape_id == target_shape.shape_id)
        assert baseline.summary.unsupported_visible_content_elements == 1

        apply_scale_text(src, out, target.element.slide_index, target.element.shape_id, 1.60)
        report = verify_output(src, out, baseline, target.element.element_id, analyzer, td / "verify")
        assert report.target_improved is True
        assert report.no_unknown_or_uncovered is False
        assert report.final_state == WorkflowState.REVIEW_REQUIRED


def test_opaque_foreground_shape_cannot_receive_verified_after_saved_artifact_verification(project_root):
    with tempfile.TemporaryDirectory() as td:
        td = Path(td)
        src = td / 'occluded-source.pptx'
        out = td / 'occluded-output.pptx'
        prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        target_shape = _add_box(slide, 'Small supported text', 1, 1, 8, 1, 12)
        _add_opaque_rectangle(slide, 1, 1, 8, 1)
        prs.save(src)

        analyzer = _analyzer(project_root)
        venue = VenueProfile(active_image_height_m=1.8, farthest_viewer_distance_m=10.8, measurement_basis="MEASURED")
        baseline = analyzer.analyze(src, venue, td / 'baseline')
        target = next(r for r in baseline.results if r.element.shape_id == target_shape.shape_id)
        assert baseline.summary.unsupported_visible_content_elements == 1

        apply_scale_text(src, out, target.element.slide_index, target.element.shape_id, 1.60)
        report = verify_output(src, out, baseline, target.element.element_id, analyzer, td / 'verify')

        assert report.no_unknown_or_uncovered is False
        assert report.final_state == WorkflowState.REVIEW_REQUIRED


def test_non_overlapping_picture_does_not_poison_supported_text_coverage(project_root):
    with tempfile.TemporaryDirectory() as td:
        td = Path(td)
        src = td / 'logo-away.pptx'
        image = td / "decorative.png"
        image.write_bytes(bytes.fromhex("89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c4890000000d49444154789c6360f8cfc0000003010100c9fe92ef0000000049454e44ae426082"))
        prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        _add_box(slide, 'Small supported text', 1, 1, 8, 1, 12)
        slide.shapes.add_picture(str(image), Inches(11), Inches(6), Inches(0.5), Inches(0.5))
        prs.save(src)

        result = _analyzer(project_root).analyze(
            src,
            VenueProfile(active_image_height_m=1.8, farthest_viewer_distance_m=10.8, measurement_basis="MEASURED"),
            td / 'analysis',
        )
        assert result.summary.unsupported_visible_content_elements == 0
