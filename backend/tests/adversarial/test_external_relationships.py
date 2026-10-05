from pathlib import Path
import tempfile
from pptx import Presentation
from pptx.util import Inches, Pt

from app.security.intake import validate_pptx


def test_external_hyperlink_is_detected():
    with tempfile.TemporaryDirectory() as td:
        path=Path(td)/'external.pptx'
        prs=Presentation(); slide=prs.slides.add_slide(prs.slide_layouts[6])
        box=slide.shapes.add_textbox(Inches(1),Inches(1),Inches(5),Inches(1))
        run=box.text_frame.paragraphs[0].add_run(); run.text='external'; run.font.size=Pt(18)
        run.hyperlink.address='https://example.com/'
        prs.save(path)
        report=validate_pptx(path)
        assert report.has_external_relationships


def test_external_relationship_with_legal_xml_whitespace_is_detected():
    import zipfile
    with tempfile.TemporaryDirectory() as td:
        td = Path(td)
        original = td / 'external-original.pptx'
        modified = td / 'external-whitespace.pptx'
        prs=Presentation(); slide=prs.slides.add_slide(prs.slide_layouts[6])
        box=slide.shapes.add_textbox(Inches(1),Inches(1),Inches(5),Inches(1))
        run=box.text_frame.paragraphs[0].add_run(); run.text='external'; run.font.size=Pt(18)
        run.hyperlink.address='https://example.com/'
        prs.save(original)

        with zipfile.ZipFile(original, 'r') as zin, zipfile.ZipFile(modified, 'w', zipfile.ZIP_DEFLATED) as zout:
            replaced = False
            for info in zin.infolist():
                data = zin.read(info.filename)
                if info.filename.endswith('.rels') and b'TargetMode="External"' in data:
                    data = data.replace(b'TargetMode="External"', b'TargetMode = "External"')
                    replaced = True
                zout.writestr(info, data)
        assert replaced
        report = validate_pptx(modified)
        assert report.has_external_relationships
