from pathlib import Path
import tempfile

import pytest
from pptx import Presentation

from app.security.intake import IntakeError, MAX_SLIDES, validate_pptx


def test_slide_count_is_bounded_before_rendering():
    with tempfile.TemporaryDirectory() as td:
        path = Path(td) / 'too-many-slides.pptx'
        prs = Presentation()
        for _ in range(MAX_SLIDES + 1):
            prs.slides.add_slide(prs.slide_layouts[6])
        prs.save(path)
        with pytest.raises(IntakeError, match='Slide count exceeds'):
            validate_pptx(path)
