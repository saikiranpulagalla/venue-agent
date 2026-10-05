from pathlib import Path
import tempfile
import copy

import pytest

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.oxml.shapes.groupshape import CT_GroupShape
from pptx.util import Inches, Pt

from app.domain.models import MutationCapability
from app.pptx.parser import parse_pptx_shapes


def _run(box, text: str, size: int = 18, font: str = "DejaVu Sans"):
    box.text_frame.clear()
    r = box.text_frame.paragraphs[0].add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.name = font


def test_text_to_fit_shape_is_review_only():
    with tempfile.TemporaryDirectory() as td:
        path = Path(td) / "autofit.pptx"
        prs = Presentation(); slide = prs.slides.add_slide(prs.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
        box.text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
        _run(box, "Autofit text")
        prs.save(path)
        rec = parse_pptx_shapes(path)[0]
    assert rec.mutation_capability == MutationCapability.REVIEW_ONLY
    assert any(n.startswith("automatic_text_fitting") for n in rec.notes)


def test_manual_prominent_top_text_is_conservatively_recognized_as_title():
    with tempfile.TemporaryDirectory() as td:
        path = Path(td) / "manual-title.pptx"
        prs = Presentation(); slide = prs.slides.add_slide(prs.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(0.3), Inches(8), Inches(0.8))
        box.text_frame.auto_size = MSO_AUTO_SIZE.NONE
        _run(box, "Quarterly Results", size=30)
        prs.save(path)
        rec = parse_pptx_shapes(path)[0]
    assert rec.role == "TITLE"


@pytest.mark.parametrize("text,top,size", [
    ("Manual 22pt heading", 2.0, 22),
    ("\u0645\u0631\u062d\u0628\u0627 \u0628\u0627\u0644\u062d\u0636\u0648\u0631", 2.0, 22),
])
def test_short_manual_heading_in_top_third_is_recognized_as_title(text, top, size):
    with tempfile.TemporaryDirectory() as td:
        path = Path(td) / "manual-heading.pptx"
        prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(top), Inches(8), Inches(0.6))
        _run(box, text, size=size)
        prs.save(path)
        rec = parse_pptx_shapes(path)[0]
    assert rec.role == "TITLE"


def test_large_body_text_is_not_broadly_reclassified_as_title():
    with tempfile.TemporaryDirectory() as td:
        path = Path(td) / "large-body.pptx"
        prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(10), Inches(2))
        _run(box, "A large body paragraph with more than fourteen words to avoid heading classification safely today", size=24)
        prs.save(path)
        rec = parse_pptx_shapes(path)[0]
    assert rec.role == "BODY"


def test_table_text_is_surfaced_as_unsupported_visible_content():
    with tempfile.TemporaryDirectory() as td:
        path = Path(td) / "table.pptx"
        prs = Presentation(); slide = prs.slides.add_slide(prs.slide_layouts[6])
        table = slide.shapes.add_table(1, 1, Inches(1), Inches(1), Inches(4), Inches(1)).table
        table.cell(0, 0).text = "Visible table text"
        prs.save(path)
        recs = parse_pptx_shapes(path)
    assert len(recs) == 1
    assert recs[0].force_not_analyzable is True
    assert recs[0].unsupported_visible_content is True
    assert "unsupported_visible_table_text" in recs[0].notes


def test_structural_text_element_count_is_bounded():
    from app.pptx.parser import MAX_STRUCTURAL_TEXT_ELEMENTS, ParserLimitError
    import pytest
    with tempfile.TemporaryDirectory() as td:
        path = Path(td) / "too-many-textboxes.pptx"
        prs = Presentation(); slide = prs.slides.add_slide(prs.slide_layouts[6])
        for i in range(MAX_STRUCTURAL_TEXT_ELEMENTS + 1):
            box = slide.shapes.add_textbox(Inches(0.1), Inches(0.1), Inches(1), Inches(0.2))
            box.text_frame.auto_size = MSO_AUTO_SIZE.NONE
            _run(box, f"x{i}", size=8)
        prs.save(path)
        with pytest.raises(ParserLimitError, match="structural_text_element_limit_exceeded"):
            parse_pptx_shapes(path)


def _opaque_rectangle(slide, left, top, width, height):
    rect = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    rect.fill.solid()
    rect.fill.fore_color.rgb = RGBColor(0, 0, 0)
    rect.line.fill.background()
    return rect


def test_foreground_opaque_shape_covering_text_is_explicit_unsupported_coverage():
    with tempfile.TemporaryDirectory() as td:
        path = Path(td) / "covered-text.pptx"
        prs = Presentation(); slide = prs.slides.add_slide(prs.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
        box.text_frame.auto_size = MSO_AUTO_SIZE.NONE
        _run(box, "Body text")
        _opaque_rectangle(slide, 1, 1, 5, 1)  # later in slide.shapes = foreground
        prs.save(path)
        recs = parse_pptx_shapes(path)

    assert len(recs) == 2
    assert any(r.unsupported_visible_content for r in recs)
    assert any("potential_text_occluder" in n for r in recs for n in r.notes)


def test_partial_foreground_occlusion_exceeding_threshold_is_unsupported_coverage():
    with tempfile.TemporaryDirectory() as td:
        path = Path(td) / "partial-covered-text.pptx"
        prs = Presentation(); slide = prs.slides.add_slide(prs.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
        box.text_frame.auto_size = MSO_AUTO_SIZE.NONE
        _run(box, "Body text")
        _opaque_rectangle(slide, 1, 1, 0.5, 1)
        prs.save(path)
        recs = parse_pptx_shapes(path)

    assert any(r.unsupported_visible_content for r in recs)


def test_background_or_non_overlapping_decoration_does_not_block_text_coverage():
    with tempfile.TemporaryDirectory() as td:
        path = Path(td) / "harmless-decoration.pptx"
        prs = Presentation(); slide = prs.slides.add_slide(prs.slide_layouts[6])
        _opaque_rectangle(slide, 0, 0, 13, 0.4)  # background / earlier z-order
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
        box.text_frame.auto_size = MSO_AUTO_SIZE.NONE
        _run(box, "Body text")
        _opaque_rectangle(slide, 10, 6, 0.5, 0.5)  # foreground but away
        prs.save(path)
        recs = parse_pptx_shapes(path)

    assert len(recs) == 1
    assert recs[0].unsupported_visible_content is False


def test_intersecting_inherited_layout_visual_is_explicit_unsupported_coverage():
    with tempfile.TemporaryDirectory() as td:
        path = Path(td) / "inherited-overlay.pptx"
        prs = Presentation()
        layout = prs.slide_layouts[6]
        staging_slide = prs.slides.add_slide(layout)
        inherited = _opaque_rectangle(staging_slide, 1, 1, 5, 1)
        # python-pptx exposes no add-shape API for layouts. Cloning a normal
        # shape element into the layout creates the same inherited OOXML case
        # the parser must represent.
        layout.shapes._spTree.insert_element_before(copy.deepcopy(inherited._element), "p:extLst")
        slide = prs.slides.add_slide(layout)
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
        box.text_frame.auto_size = MSO_AUTO_SIZE.NONE
        _run(box, "Body text")
        prs.save(path)
        recs = parse_pptx_shapes(path)

    assert any(r.unsupported_visible_content for r in recs)
    assert any("inherited_layout" in n for r in recs for n in r.notes)


def test_visual_only_group_covering_text_is_explicit_unsupported_coverage():
    with tempfile.TemporaryDirectory() as td:
        path = Path(td) / "group-overlay.pptx"
        prs = Presentation(); slide = prs.slides.add_slide(prs.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
        box.text_frame.auto_size = MSO_AUTO_SIZE.NONE
        _run(box, "Body text")
        child = _opaque_rectangle(slide, 1, 1, 5, 1)
        group = CT_GroupShape.new_grpSp(99, "visual-only-group")
        xfrm = group.grpSpPr.xfrm
        xfrm.off.x = Inches(1); xfrm.off.y = Inches(1)
        xfrm.ext.cx = Inches(5); xfrm.ext.cy = Inches(1)
        xfrm.chOff.x = Inches(1); xfrm.chOff.y = Inches(1)
        xfrm.chExt.cx = Inches(5); xfrm.chExt.cy = Inches(1)
        group.insert_element_before(copy.deepcopy(child._element), "p:extLst")
        slide.shapes._spTree.remove(child._element)
        slide.shapes._spTree.insert_element_before(group, "p:extLst")
        prs.save(path)
        recs = parse_pptx_shapes(path)

    assert any(r.unsupported_visible_content for r in recs)
    assert any("potential_text_occluder" in n for r in recs for n in r.notes)
