from __future__ import annotations

from pathlib import Path

import fitz
from pptx import Presentation

from app.domain.models import AnalysisResult, VerificationReport, WorkflowState
from app.pptx.renderer import render_pptx_to_pdf
from app.repair.transition import audit_analysis_transition
from app.security.intake import sha256_path
from app.venue.analyzer import VenueAnalyzer


def _normalized_visible_text(path: Path) -> list[str]:
    prs = Presentation(str(path))
    texts: list[str] = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and (shape.text or "").strip():
                texts.append(" ".join((shape.text or "").split()))
            elif getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    for cell in row.cells:
                        if (cell.text or "").strip():
                            texts.append(" ".join((cell.text or "").split()))
    return texts


def _rendered_text(path: Path, work: Path) -> str:
    pdf = render_pptx_to_pdf(path, work)
    doc = fitz.open(pdf)
    return " ".join(" ".join(page.get_text().split()) for page in doc)


def _find_result(analysis: AnalysisResult, element_id: str):
    return next((r for r in analysis.results if r.element.element_id == element_id), None)


def _coverage(analysis: AnalysisResult) -> float:
    return analysis.summary.target_coverage


def verify_output(
    source: Path,
    output: Path,
    baseline: AnalysisResult,
    target_element_ids: list[str] | str,
    analyzer: VenueAnalyzer,
    work_dir: Path,
) -> VerificationReport:
    reasons: list[str] = []
    if isinstance(target_element_ids, str):
        target_element_ids = [target_element_ids]
    source_sha_ok = sha256_path(source) == baseline.source_sha256
    output_reopens = True
    try:
        Presentation(str(output))
    except Exception:
        output_reopens = False
        reasons.append("output_does_not_reopen")

    protected_text_unchanged = False
    target_improved = False
    target_improvements: dict[str, bool] = {}
    analysis_universe_preserved = False
    layout_safe = False
    no_unknown_or_uncovered = False
    after = None
    if output_reopens:
        try:
            protected_text_unchanged = _normalized_visible_text(source) == _normalized_visible_text(output)
            if not protected_text_unchanged:
                reasons.append("pptx_visible_text_changed")

            # Fresh rendered-text equality catches text loss, while the
            # transition audit below checks where that text ended up.
            src_rendered = _rendered_text(source, work_dir / "source-render")
            out_rendered = _rendered_text(output, work_dir / "output-render")
            if src_rendered != out_rendered:
                protected_text_unchanged = False
                reasons.append("rendered_text_changed_or_lost")

            after = analyzer.analyze(output, baseline.venue, work_dir / "verify-analysis")
            transition = audit_analysis_transition(baseline, after, target_element_ids)
            no_unknown_or_uncovered = not after.summary.has_unknown_or_uncovered
            if not no_unknown_or_uncovered:
                reasons.append("unknown_or_uncovered_content_remains")
            analysis_universe_preserved = transition.analysis_universe_preserved
            layout_safe = transition.layout_safe
            reasons.extend(transition.reasons)
            if not analysis_universe_preserved:
                reasons.append("analysis_universe_not_preserved")
            if not layout_safe:
                reasons.append("layout_regression")

            for element_id in target_element_ids:
                b = _find_result(baseline, element_id)
                a = _find_result(after, element_id)
                improved = bool(b and a and (a.actual_percent or 0) > (b.actual_percent or 0))
                target_improvements[element_id] = improved
                if not improved:
                    reasons.append(f"target_not_improved:{element_id}")
            target_improved = bool(target_element_ids) and all(target_improvements.values())
            if after.summary.below_target_elements > baseline.summary.below_target_elements:
                reasons.append("new_supported_violation")
                target_improved = False
        except Exception as e:
            reasons.append(f"verification_failed:{type(e).__name__}")

    if not source_sha_ok:
        reasons.append("source_hash_changed")

    reasons = sorted(set(reasons))
    final = WorkflowState.VERIFIED if all([
        source_sha_ok,
        output_reopens,
        protected_text_unchanged,
        target_improved,
        analysis_universe_preserved,
        layout_safe,
        no_unknown_or_uncovered,
    ]) and "new_supported_violation" not in reasons else WorkflowState.REVIEW_REQUIRED

    return VerificationReport(
        output_sha256=sha256_path(output) if output.exists() else "",
        source_sha256_unchanged=source_sha_ok,
        protected_text_unchanged=protected_text_unchanged,
        output_reopens=output_reopens,
        target_improved=target_improved,
        analysis_universe_preserved=analysis_universe_preserved,
        layout_safe=layout_safe,
        no_unknown_or_uncovered=no_unknown_or_uncovered,
        final_state=final,
        reasons=reasons,
        target_improvements=target_improvements,
        target_coverage_before=_coverage(baseline),
        target_coverage_after=_coverage(after) if after is not None else None,
    )
