from __future__ import annotations

from hashlib import sha256
import json
from pathlib import Path

from pptx import Presentation
import fitz


def normalized_visible_text(path: Path) -> list[str]:
    prs = Presentation(str(path))
    texts: list[str] = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and (shape.text or "").strip():
                texts.append(" ".join((shape.text or "").split()))
    return texts


def pptx_reopens(path: Path) -> bool:
    try:
        Presentation(str(path))
        return True
    except Exception:
        return False


def canonical_fingerprint(payload: object) -> str:
    raw = json.dumps(payload, sort_keys=True, separators=(",", ":"), default=str).encode("utf-8")
    return sha256(raw).hexdigest()


def normalized_rendered_pdf_text(pdf_path: Path) -> str:
    """Return normalized visible text from a rendered PDF artifact.

    Used as an independent invariant against clipping/loss introduced by a
    candidate plan. This deliberately does not reuse PPTX XML text.
    """
    doc = fitz.open(pdf_path)
    try:
        return " ".join(" ".join(page.get_text().split()) for page in doc)
    finally:
        doc.close()
